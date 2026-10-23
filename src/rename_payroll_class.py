import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# Project root
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

sys.path.insert(
    0,
    str(BASE_DIR),
)

from config import CLASS_RENAMES, OUTPUT_DATA_DIR, PROCESSED_DATA_DIR


# ============================================================
# One-off class rename in existing data files (Strategy)
#
# "Μισθοδοσία συνδεδεμένη με το επενδυτικό σχέδιο (όπως
# προβλέπεται από τον ΓΑΚ)" → "Μισθοδοσία συνδεδεμένη με το
# επενδυτικό σχέδιο", wherever the old name appears (expense
# class, subcategory, or texts built from them).
#
# The preparation scripts apply the same rename from now on.
# Small-POC files stay as they are, so the baseline and the
# fine-tuned model are compared on the same texts.
# data/raw is never touched.
#
# Default: dry run (counts only). --apply writes the files.
# ============================================================

SMALL_POC_FILES = {
    "training_pairs_poc.csv",
    "training_pairs_poc.xlsx",
    "embeddings_on_pairs_results.xlsx",
    "poc_annotation_sample.xlsx",
}

SMALL_POC_PREFIXES = (
    "review_100_",
    "review_300_",
)


def is_small_poc_file(path):

    return (
        path.name in SMALL_POC_FILES
        or path.name.startswith(SMALL_POC_PREFIXES)
    )


def target_files():

    files = []

    for folder in [PROCESSED_DATA_DIR, OUTPUT_DATA_DIR]:
        for path in sorted(folder.iterdir()):
            if path.suffix in {".csv", ".xlsx"} and not is_small_poc_file(path):
                files.append(path)

    return files


def rename_text(value):

    for old, new in CLASS_RENAMES.items():
        value = value.replace(old, new)

    return value


def count_old(text):

    return sum(text.count(old) for old in CLASS_RENAMES)


# ============================================================
# CSV: plain text replace, encoding and layout unchanged
# ============================================================

def process_csv(path, apply):

    text = path.read_text(encoding="utf-8-sig")
    found = count_old(text)

    if apply and found:
        path.write_text(rename_text(text), encoding="utf-8-sig")

    return found


# ============================================================
# Excel: cell-by-cell replace, formatting and sheets unchanged
# ============================================================

def process_xlsx(path, apply):

    workbook = load_workbook(path)
    found = 0

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    n = count_old(cell.value)
                    if n:
                        found += n
                        cell.value = rename_text(cell.value)

    if apply and found:
        workbook.save(path)

    return found


# ============================================================
# Main
# ============================================================

def main() -> None:

    parser = argparse.ArgumentParser(
        description="Rename the payroll expense class in existing data files."
    )
    parser.add_argument("--apply", action="store_true", help="Write the files (default: dry run).")
    args = parser.parse_args()

    total = 0

    for path in target_files():
        if path.suffix == ".csv":
            found = process_csv(path, args.apply)
        else:
            found = process_xlsx(path, args.apply)
        if found:
            total += found
            print(f"{found:>5}  {path.relative_to(BASE_DIR)}")

    print("=" * 60)
    print("Mode:", "APPLY" if args.apply else "DRY RUN (nothing written)")
    print("Occurrences of the old class name:", total)

    if args.apply:
        remaining = sum(
            process_csv(path, False) if path.suffix == ".csv" else process_xlsx(path, False)
            for path in target_files()
        )
        assert remaining == 0, f"{remaining} occurrences remain after the rename."
        print("Check: no occurrence of the old name remains in the target files.")


if __name__ == "__main__":
    main()
